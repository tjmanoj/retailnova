"""Build the RetailNova prototype deck on top of the official AI Builder Cup template.

Usage:
    .venv/bin/python docs/submission/build_deck.py [--template PATH] [--links PATH]

Reads measured numbers from evals/RESULTS.json and screenshots from docs/screenshots/.
Writes docs/submission/RetailNova-Prototype-Deck.pptx (convert to PDF with LibreOffice).
"""
from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = REPO.parent / "dashboard" / "AI-Builder-Cup-Prototype-Template.pptx"
OUT = REPO / "docs" / "submission" / "RetailNova-Prototype-Deck.pptx"
SHOTS = REPO / "docs" / "screenshots"
RESULTS = REPO / "evals" / "RESULTS.json"

FONT = "Arial"
INK = RGBColor(0x1A, 0x1A, 0x1F)
MUTED = RGBColor(0x5B, 0x5B, 0x66)
MAROON = RGBColor(0x7A, 0x1E, 0x2C)
MARIGOLD = RGBColor(0xE8, 0x9A, 0x0C)
GREEN = RGBColor(0x1E, 0x7A, 0x4C)
CARD = RGBColor(0xF6, 0xF3, 0xEE)
LINE = RGBColor(0xDD, 0xD6, 0xCC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Content area under the template's header band and slide title.
X0, Y0, W, H = 0.32, 1.30, 9.36, 4.12


# ---------- primitives ----------

def _style_run(run, size, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, paras, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """paras: list of str or (str, dict(size, bold, color, italic, bullet))."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.04))
    for i, p in enumerate(paras):
        content, opts = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        if opts.get("space_after") is not None:
            para.space_after = Pt(opts["space_after"])
        prefix = "•  " if opts.get("bullet") else ""
        run = para.add_run()
        run.text = prefix + content
        _style_run(run, opts.get("size", size), opts.get("bold", bold), opts.get("color", color), opts.get("italic", False))
    return box


def rect(slide, x, y, w, h, fill=CARD, line=None, radius=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def card(slide, x, y, w, h, title, body, accent=MAROON, title_size=11.5, body_size=9.5):
    rect(slide, x, y, w, h)
    rect(slide, x, y, 0.06, h, fill=accent, radius=False)
    text(slide, x + 0.14, y + 0.07, w - 0.2, 0.34, [(title, {"bold": True, "size": title_size, "color": accent})])
    body_paras = body if isinstance(body, list) else [body]
    text(slide, x + 0.14, y + 0.40, w - 0.22, h - 0.45,
         [(b, {"size": body_size, "color": INK, "space_after": 2}) for b in body_paras], spacing=1.08)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.75):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return conn


def picture(slide, path: Path, x, y, w, h, border=True):
    """Fit an image inside the box, preserving aspect ratio, centered."""
    if not path.exists():
        rect(slide, x, y, w, h, fill=CARD, line=LINE)
        text(slide, x, y + h / 2 - 0.15, w, 0.3, [(f"[{path.name}]", {"size": 9, "color": MUTED})], align=PP_ALIGN.CENTER)
        return None
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    pic = slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic


def table(slide, x, y, w, rows, col_w, size=9.5, header_fill=MAROON, row_h=0.28):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (CARD if i % 2 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run()
            run.text = str(val)
            _style_run(run, size, bold=(i == 0), color=WHITE if i == 0 else INK)
    return tbl


def set_title(slide, new_text=None, size=16):
    """Restyle the template's title box (keeps its position) and optionally replace its text."""
    for shp in slide.shapes:
        if shp.has_text_frame and abs(shp.top - Inches(0.73)) < Inches(0.05):
            tf = shp.text_frame
            if new_text is not None:
                for p in list(tf.paragraphs)[1:]:
                    p._p.getparent().remove(p._p)
                tf.paragraphs[0].text = ""
                run = tf.paragraphs[0].add_run()
                run.text = new_text
            for p in tf.paragraphs:
                for r in p.runs:
                    _style_run(r, size, bold=True, color=INK)
            shp.height = Inches(0.45)
            return shp
    return None


def delete_slide(prs, index):
    sld_ids = prs.slides._sldIdLst
    rid = sld_ids[index].get(qn("r:id"))
    prs.part.drop_rel(rid)
    sld_ids.remove(sld_ids[index])


# ---------- numbers ----------

def load_results() -> dict:
    if RESULTS.exists():
        return json.loads(RESULTS.read_text())
    return {}


def fmt(v, kind="num", missing="TBD"):
    if v is None:
        return missing
    if kind == "pct":
        return f"{v:.1f}%"
    if kind == "inr":
        return f"₹{v:,.0f}"
    if kind == "lakh":
        return f"₹{v / 1e5:,.1f} lakh"
    if kind == "s":
        return f"{v:.1f}s"
    return f"{v:,.0f}" if isinstance(v, (int, float)) and abs(v) >= 100 else f"{v}"


# ---------- slides ----------

def build(template: Path, links: dict):
    prs = Presentation(str(template))
    r = load_results()
    bt = r.get("backtest", {})
    ev = r.get("concierge", {})
    lat = r.get("latency", {})
    cost = r.get("cost", {})

    # Remove the template's instruction note slide (index 1).
    delete_slide(prs, 1)
    s = prs.slides

    # 1. Team details (cover)
    cover = s[0]
    for shp in cover.shapes:
        if shp.has_text_frame and "Team Details" in shp.text_frame.text:
            values = {"Team name": "RetailNova", "Team leader name": "Suvetha S",
                      "Problem Statement": "Retail & Commerce: Intelligent Customer and Business Experiences"}
            for p in shp.text_frame.paragraphs:
                for r_ in p.runs:
                    _style_run(r_, r_.font.size.pt if r_.font.size else 16, bold=r_.font.bold or False)
                key = p.text.strip().rstrip(":").strip()
                if key in values:
                    run = p.add_run()
                    run.text = " " + values[key]
                    _style_run(run, 14, bold=True, color=MAROON)
            shp.width = Inches(5.6)
    text(cover, 6.05, 3.25, 3.75, 2.2, [
        ("RetailNova", {"size": 24, "bold": True, "color": MAROON}),
        ("Every conversation is a demand signal.", {"size": 13, "bold": True, "color": INK, "space_after": 6}),
        ("Gemini agents on Google Cloud that turn what shoppers ask for into what stores stock.",
         {"size": 10.5, "color": MUTED}),
        ("Team: Suvetha S (lead) · Manoj T", {"size": 9.5, "color": MUTED}),
    ])

    # 2. Brief about the idea
    sl = s[1]
    set_title(sl, "Brief about the idea")
    text(sl, X0, Y0, 4.55, 0.32, [("The problem: stockouts hide demand", {"bold": True, "size": 13, "color": MAROON})])
    text(sl, X0, Y0 + 0.36, 4.55, 1.4, [
        ("Retailers forecast from what they sold, not from what customers wanted. When an item is out of stock, "
         "its demand disappears from the data: zero sales → a smaller forecast → a smaller order → the gap repeats "
         "every season.", {"size": 10.5}),
        ("The shopper who asked for a maroon silk saree before Diwali and walked away is never counted.",
         {"size": 10.5, "italic": True, "color": MUTED}),
    ])
    stats = [("$1.2T", "lost to out-of-stocks globally each year", "IHL Group, 2025"),
             ("37%", "of shoppers go to another store when an item is unavailable", "IGD ShopperVista"),
             ("₹5.4L cr", "Navratri → Diwali 2025 trade; ~85% via physical retail", "CAIT, Oct 2025")]
    for i, (big, label, src) in enumerate(stats):
        cx = X0 + i * 1.53
        rect(sl, cx, Y0 + 1.95, 1.45, 1.55)
        text(sl, cx + 0.05, Y0 + 2.02, 1.35, 0.45, [(big, {"bold": True, "size": 18, "color": MAROON})])
        text(sl, cx + 0.05, Y0 + 2.5, 1.35, 0.75, [(label, {"size": 8.5})])
        text(sl, cx + 0.05, Y0 + 3.2, 1.35, 0.25, [(src, {"size": 7, "color": MUTED, "italic": True})])
    rect(sl, 5.05, Y0, 4.63, 3.5, fill=WHITE, line=LINE)
    text(sl, 5.2, Y0 + 0.08, 4.4, 0.32, [("RetailNova closes the loop", {"bold": True, "size": 13, "color": GREEN})])
    loop = [("Ask", "Shoppers ask in their language by text, photo or voice note; staff log walk-in requests."),
            ("Serve", "A Gemini concierge answers only from the real catalog and live stock in every store."),
            ("Sense", "Each conversation becomes a structured demand signal: what, where, price, and why it was unmet."),
            ("Plan", "BigQuery TimesFM forecasts plus signals → an agent proposes transfers and reorders, with evidence."),
            ("Act", "The manager approves; inventory moves; the waiting shopper is notified.")]
    for i, (k, v) in enumerate(loop):
        yy = Y0 + 0.5 + i * 0.58
        rect(sl, 5.2, yy, 0.72, 0.46, fill=MAROON if i % 2 == 0 else MARIGOLD)
        text(sl, 5.2, yy + 0.08, 0.72, 0.3, [(k, {"bold": True, "size": 10, "color": WHITE})], align=PP_ALIGN.CENTER)
        text(sl, 6.0, yy + 0.02, 3.6, 0.5, [(v, {"size": 9.5})])
    text(sl, X0, Y0 + 3.62, 9.3, 0.4, [("Built for Indian festive retail: physical stores + online, 21 days before Diwali (8 Nov 2026). "
                                         "All demo data is synthetic; Nova Bazaar is a fictional retailer.",
                                         {"size": 8.5, "color": MUTED, "italic": True})])

    # 3. Opportunities
    sl = s[2]
    set_title(sl, "Opportunities")
    cols = [
        ("How is it different?", [
            "Shopping assistants stop at the conversation. Forecasting tools only see sales.",
            "RetailNova connects them: demand that customers voice becomes a planning input the same day.",
            "Google's Conversational Commerce agent serves the shopper side; RetailNova closes the loop to store operations."]),
        ("How does it solve the problem?", [
            "Captures unmet demand when it is expressed: chat, photo search, store assist.",
            "Turns it into structured signals, then a demand-sensed TimesFM forecast.",
            "An agent proposes evidence-backed transfers and reorders; a human approves; the shopper is notified."]),
        ("USP", [
            "Demand you can't see in sales, made visible and actionable.",
            f"Measured (synthetic backtest): forecast error {fmt(bt.get('wape_sales_only'), 'pct')} → "
            f"{fmt(bt.get('wape_demand_sensed'), 'pct')}; lost revenue cut by {fmt(bt.get('lost_revenue_reduction_pct'), 'pct')}.",
            f"Grounded: {fmt(ev.get('grounding_rate'), 'pct')} of recommended products verified against catalog and stock."]),
    ]
    for i, (t, body) in enumerate(cols):
        card(sl, X0 + i * 3.14, Y0 + 0.05, 3.02, 3.95, t, body, accent=[MAROON, GREEN, MARIGOLD][i], title_size=12.5, body_size=11.5)

    # 4. Features
    sl = s[3]
    set_title(sl, "List of features offered by the solution")
    feats = [
        ("Multilingual concierge", "Text, photo or voice note in Tamil, Hindi or English. Grounded in catalog + live stock; never invents products or prices."),
        ("Nova remembers you", "Agent Platform Memory Bank recalls taste and budget across visits. One-tap “Forget me”."),
        ("Omnichannel availability", "Stock across 6 stores + online. Reserve with a QR pickup pass or tap “Notify me”."),
        ("Every ask is a signal", "Chats, voice notes, photos and walk-in requests logged by staff become structured demand signals."),
        ("Demand Radar", "Sales-only vs demand-sensed TimesFM forecasts, assortment gaps and revenue protected."),
        ("Agentic action plan", "Parallel steps incl. live Google Search market trends → transfers and reorders with evidence."),
        ("Approve → notify", "One click moves stock and notifies waiting shoppers. Morning brief for the manager."),
        ("Guardrails", "Model Armor screens every prompt; server-side grounding; rate limits; no personal data."),
    ]
    for i, (t, b) in enumerate(feats):
        cx, cy = X0 + (i % 4) * 2.35, Y0 + 0.05 + (i // 4) * 2.02
        card(sl, cx, cy, 2.25, 1.92, t, b, accent=[MAROON, MARIGOLD, GREEN, MAROON][i % 4], title_size=11.5, body_size=10.5)

    # 5. Process flow
    sl = s[4]
    set_title(sl, "Process flow diagram")
    steps = [("1 · Shopper asks", "Text, photo or voice note in any language; staff log walk-ins"),
             ("2 · Concierge", "Model Armor check → Gemini 3.6 Flash + tools: vector search, stock, memory"),
             ("3 · Signal", "Flash-Lite extracts a DemandSignal → Firestore (live)"),
             ("4 · Plan", "TimesFM + signals + inventory + Google Search trends → proposals"),
             ("5 · Approve", "Manager approves → inventory moves → shopper notified")]
    bw, gap = 1.7, 0.215
    for i, (t, b) in enumerate(steps):
        bx = X0 + i * (bw + gap)
        card(sl, bx, Y0 + 0.35, bw, 1.55, t, b, accent=[MAROON, MARIGOLD, GREEN, MARIGOLD, MAROON][i], title_size=10.5, body_size=9)
        if i < len(steps) - 1:
            arrow(sl, bx + bw + 0.02, Y0 + 1.12, bx + bw + gap - 0.02, Y0 + 1.12)
    arrow(sl, X0 + 4 * (bw + gap) + bw / 2, Y0 + 1.95, X0 + 4 * (bw + gap) + bw / 2, Y0 + 2.35, color=GREEN)
    arrow(sl, X0 + 4 * (bw + gap) + bw / 2, Y0 + 2.35, X0 + bw / 2, Y0 + 2.35, color=GREEN)
    arrow(sl, X0 + bw / 2, Y0 + 2.35, X0 + bw / 2, Y0 + 1.95, color=GREEN)
    text(sl, X0 + 2.2, Y0 + 2.4, 5.0, 0.3, [("“Reserved for you at T. Nagar”: the loop closes", {"size": 9.5, "italic": True, "color": GREEN})],
         align=PP_ALIGN.CENTER)
    rect(sl, X0, Y0 + 2.9, W, 1.1, fill=CARD)
    text(sl, X0 + 0.15, Y0 + 2.98, W - 0.3, 1.0, [
        ("Use case (hero flow)", {"bold": True, "size": 10.5, "color": MAROON}),
        ("Priya asks for a maroon silk saree under ₹5,000 for Diwali pickup at T. Nagar. It is out of stock there but 14 sit in Coimbatore, "
         "so she taps “Notify me”. Arjun's Radar shows the unmet-demand spike live; the plan proposes a transfer from Coimbatore plus a "
         "reorder; he approves; Priya's chat shows it is reserved for her.", {"size": 9.5}),
    ])

    # 6. Wireframes / mocks
    sl = s[5]
    set_title(sl, "Wireframes / mock diagrams of the solution")
    picture(sl, SHOTS / "shop-mobile.png", X0, Y0, 2.2, 4.05)
    picture(sl, SHOTS / "manager.png", X0 + 2.35, Y0, 7.0, 4.05)

    # 7. Architecture
    sl = s[6]
    set_title(sl, "Architecture diagram of the proposed solution")
    arch_png = REPO / "docs" / "architecture.png"
    if arch_png.exists():
        picture(sl, arch_png, X0, Y0, W, 4.1, border=False)
    else:
        draw_architecture(sl)

    # 8. Technologies
    sl = s[7]
    set_title(sl, "Technologies used in the solution")
    rows = [["Google service", "Used for"],
            ["Gemini 3.6 Flash (Agent Platform)", "Concierge reasoning + tool calling, voice-note and image understanding, multilingual chat, ops planning, morning brief"],
            ["Gemini 3.5 Flash-Lite", "Structured DemandSignal extraction (JSON schema); hedged fallback model"],
            ["Google Search grounding", "Live festive market trends with cited sources inside the action plan"],
            ["Agent Platform Memory Bank", "Remembers shopper preferences across visits (preferences only); “Forget me”"],
            ["Model Armor", "Screens every prompt for injection, jailbreak and harmful content before tools run"],
            ["Gemini Embedding 2 · Gemini 3.1 Flash Image", "768-dim vectors for semantic + photo search; generated the catalog photo set"],
            ["Agent Development Kit (ADK)", "Agents, function tools, sessions, streamed events"],
            ["Cloud Run", "Serverless agent API (FastAPI); autoscaling, keyless service identity"],
            ["Firebase Hosting", "Global CDN for the UI (installable PWA); /api rewrite to Cloud Run"],
            ["Firestore", "Catalog, inventory, signals, plans; vector search; real-time listeners in the UI"],
            ["BigQuery + AI.FORECAST (TimesFM)", "2-year sales history; forecasts with no model training; backtest analytics"],
            ["Cloud Trace · Monitoring · Scheduler · Build", "Spans per agent/tool call, uptime alerts, nightly demo reset, container builds"]]
    table(sl, X0, Y0 - 0.05, W, rows, [3.1, 6.26], size=8.5, row_h=0.315)

    # 9. Estimated cost
    sl = s[8]
    set_title(sl, "Estimated implementation cost")
    rows = [["Item", "Prototype to date (est.)", "Pilot: 50 stores, 100k chats/month (est.)"],
            ["Gemini (agents + extraction)", fmt(cost.get("proto_gemini_inr"), "inr"), fmt(cost.get("pilot_gemini_inr"), "inr") + " / month"],
            ["Cloud Run + Firebase Hosting", fmt(cost.get("proto_run_inr"), "inr"), fmt(cost.get("pilot_run_inr"), "inr") + " / month"],
            ["Firestore + BigQuery", fmt(cost.get("proto_data_inr"), "inr"), fmt(cost.get("pilot_data_inr"), "inr") + " / month"],
            ["Total", fmt(cost.get("proto_total_inr"), "inr"), fmt(cost.get("pilot_total_inr"), "inr") + " / month"]]
    table(sl, X0, Y0, W, rows, [3.2, 2.6, 3.56], size=10, row_h=0.42)
    text(sl, X0, Y0 + 2.3, W, 1.7, [
        (f"Measured cost per 1,000 shopper conversations: {fmt(cost.get('per_1k_chats_inr'), 'inr')} "
         f"(from real token usage on Gemini 3.6 Flash + Flash-Lite).", {"size": 10.5, "bold": True}),
        ("Serverless by design: Cloud Run scales to zero, Firestore and BigQuery bill per use, and no GPUs are needed "
         "(TimesFM runs inside BigQuery).", {"size": 10}),
        ("Payback logic: recovering even 1% of festive lost sales for a 50-store chain outweighs the platform cost "
         "(see the benchmark slide for simulated recovery).", {"size": 10, "color": MUTED}),
    ])

    # 10. Snapshots
    sl = s[9]
    set_title(sl, "Snapshots of the prototype")
    pairs = [("live-loop.png", "1 · Priya asks in Hinglish → grounded cards, sold out at T. Nagar → the demand signal lands live on Arjun's radar"),
             ("notify.png", "2 · Arjun approves the plan: 12 units move from RS Puram → Priya's phone: “Reserved for you at T. Nagar”")]
    for i, (name, cap) in enumerate(pairs):
        cx = X0 + i * 4.73
        picture(sl, SHOTS / name, cx, Y0, 4.6, 2.9)
        text(sl, cx, Y0 + 2.98, 4.6, 0.6, [(cap, {"size": 9.5, "color": INK})])
    text(sl, X0, Y0 + 3.62, W, 0.4, [("Live at https://retailnova-ai.web.app (no login). Also: mobile shopper chat, photo search, Tamil, manager Q&A.",
                                      {"size": 9, "color": MUTED, "italic": True})])

    # 11. Performance / benchmarking
    sl = s[10]
    set_title(sl, "Prototype performance report / benchmarking")
    text(sl, X0, Y0 - 0.02, 5.4, 0.3, [("Backtest: Diwali 2025 window (synthetic, true demand known)", {"bold": True, "size": 11, "color": MAROON})])
    rows = [["Metric", "Sales-only forecast", "Demand-sensed (+ signals)"],
            ["Forecast error (WAPE)", fmt(bt.get("wape_sales_only"), "pct"), fmt(bt.get("wape_demand_sensed"), "pct")],
            ["Stockout days", fmt(bt.get("stockout_days_sales_only")), fmt(bt.get("stockout_days_demand_sensed"))],
            ["Lost revenue (6 weeks)", fmt(bt.get("lost_revenue_sales_only_inr"), "lakh"),
             f'{fmt(bt.get("lost_revenue_demand_sensed_inr"), "lakh")} (−{fmt(bt.get("lost_revenue_reduction_pct"), "pct")})']]
    table(sl, X0, Y0 + 0.32, 5.4, rows, [1.9, 1.75, 1.75], size=9.5, row_h=0.36)
    text(sl, X0, Y0 + 1.85, 5.4, 0.9, [
        (bt.get("sensitivity_note", "Sensitivity across signal capture rates of 10–50% is in evals/backtest_report.md."), {"size": 9}),
        ("Method: TimesFM (BigQuery AI.FORECAST), 42-day horizon, cutoff 6 weeks before Diwali; the same replenishment policy "
         "for both arms, scored against generator ground truth.", {"size": 8.5, "color": MUTED}),
    ])
    text(sl, 5.95, Y0 - 0.02, 3.7, 0.3, [("Concierge and system", {"bold": True, "size": 11, "color": GREEN})])
    rows = [["Measure", "Result"],
            ["Grounded products (no invented SKUs)", fmt(ev.get("grounding_rate"), "pct")],
            ["Budget / size constraints respected", fmt(ev.get("constraint_adherence"), "pct")],
            ["Off-topic / injection refused", fmt(ev.get("refusal_accuracy"), "pct")],
            ["Honest out-of-stock answers", fmt(ev.get("oos_honesty"), "pct")],
            ["Signal extraction accuracy", fmt(ev.get("signal_accuracy"), "pct")],
            ["Chat latency p50 / p95", f"{fmt(lat.get('chat_p50_s'), 's')} / {fmt(lat.get('chat_p95_s'), 's')}"],
            ["Action plan (5 parallel steps incl. Search trends)", fmt(lat.get('plan_s'), 's')]]
    table(sl, 5.95, Y0 + 0.32, 3.73, rows, [2.45, 1.28], size=9, row_h=0.32)
    text(sl, X0, Y0 + 3.55, W, 0.45, [(f"Eval set: {ev.get('n_cases', 'TBD')} cases, {fmt(ev.get('case_pass_rate'), 'pct')} passing (English, Hinglish, Tamil, image, budget/size, out-of-stock, "
                                       "off-topic, prompt injection), run against the deployed API. Numbers are reported as measured.",
                                       {"size": 8.5, "color": MUTED, "italic": True})])

    # 12. Future development
    sl = s[11]
    set_title(sl, "Additional details / future development")
    card(sl, X0, Y0 + 0.05, 4.6, 3.95, "Roadmap", [
        "Real-time voice conversations with Gemini Live in regional languages",
        "Personalised festive offers from remembered preferences (consented)",
        "Checkout via the Universal Commerce Protocol",
        "Agent-to-agent supplier negotiation for reorders",
        "WhatsApp channel and in-store associate app (store-assist signals)",
        "Shelf photo audits with Gemini vision to catch phantom stock",
    ], accent=MAROON, title_size=12.5, body_size=11.5)
    card(sl, X0 + 4.75, Y0 + 0.05, 4.6, 3.95, "Responsible AI and path to production", [
        "Privacy: signals are aggregated; no personal data in analytics; consent flows aligned with India's DPDP Act 2023",
        "Human in control: no inventory action without manager approval",
        "Pilot: 5 stores × 6 weeks before a festival; success = lower stockout days and recovered lost sales vs control stores",
        "Scale: Pub/Sub signal ingestion, scheduled TimesFM runs for thousands of series, Agent Runtime for long-running agents",
    ], accent=GREEN, title_size=12.5, body_size=11.5)

    # 13. Links
    sl = s[12]
    set_title(sl, "Links")
    items = [("Live prototype (no login)", links.get("app", "https://retailnova-ai.web.app")),
             ("GitHub public repository", links.get("github", "https://github.com/tjmanoj/retailnova")),
             ("Demo video", links.get("video", "TBD"))]
    for i, (k, v) in enumerate(items):
        yy = Y0 + 0.15 + i * 0.85
        rect(sl, X0, yy, W, 0.7)
        text(sl, X0 + 0.2, yy + 0.08, 3.0, 0.5, [(k, {"bold": True, "size": 12, "color": MAROON})], anchor=MSO_ANCHOR.MIDDLE)
        box = text(sl, X0 + 3.3, yy + 0.08, 5.9, 0.5, [(v, {"size": 12, "color": INK})], anchor=MSO_ANCHOR.MIDDLE)
        if v.startswith("http"):
            box.text_frame.paragraphs[0].runs[0].hyperlink.address = v
    text(sl, X0, Y0 + 2.8, W, 0.9, [
        ("Try in 2 minutes: open the live link → “Live Loop” → tap the sample prompt → watch the signal appear → "
         "“Generate Diwali action plan” → Approve.", {"size": 10.5}),
        ("All data is synthetic; Nova Bazaar is a fictional retailer.", {"size": 9, "color": MUTED, "italic": True}),
    ])

    # 14. Extra content slide: impact and feasibility
    sl = s[13]
    text(sl, 0.29, 0.73, 9.4, 0.45, [("Impact and feasibility", {"bold": True, "size": 16})])
    card(sl, X0, Y0 + 0.05, 3.02, 3.95, "For shoppers", [
        "Answers in their own language, grounded in real stock",
        "No dead ends: other store, alternative, reserve or notify",
        "Gets the item they actually wanted, in time for the festival"], accent=MAROON, title_size=12.5, body_size=11.5)
    card(sl, X0 + 3.17, Y0 + 0.05, 3.02, 3.95, "For retailers", [
        "Sees demand that sales data hides",
        "Moves stock before the festival peak, not after",
        "Every proposal is explainable: forecast, signals, inventory as evidence"], accent=GREEN, title_size=12.5, body_size=11.5)
    card(sl, X0 + 6.34, Y0 + 0.05, 3.02, 3.95, "Why it is feasible now", [
        "Runs entirely on managed Google Cloud services",
        "No model training: Gemini + TimesFM via BigQuery",
        "Deployed and live today; measured cost and latency"], accent=MARIGOLD, title_size=12.5, body_size=11.5)

    # Move the extra content slide (impact) before the Links slide.
    ids = prs.slides._sldIdLst
    impact = ids[13]
    ids.remove(impact)
    ids.insert(12, impact)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"saved {OUT}")


def draw_architecture(sl):
    """Architecture diagram drawn with native shapes."""
    def node(x, y, w, h, title, sub, accent, ts=10, bs=8.5):
        card(sl, x, y, w, h, title, sub, accent=accent, title_size=ts, body_size=bs)
    node(X0, Y0 + 0.75, 1.55, 1.15, "Users", "Shoppers · staff · managers · judges (browser, phone, PWA)", MUTED)
    node(X0 + 1.85, Y0 + 0.75, 1.85, 1.15, "Firebase Hosting", "React UI on CDN; /api → Cloud Run; Firestore live listeners", MARIGOLD)
    node(X0 + 4.0, Y0 + 0.3, 2.45, 2.05, "Cloud Run · ADK agents", [
        "Concierge (text · photo · voice)", "Signal extractor · Store assist", "Ops planner: parallel steps → Gemini", "Manager Q&A · morning brief"], MAROON)
    node(X0 + 6.8, Y0 + 0.0, 2.55, 0.95, "Gemini on Agent Platform", "3.6 Flash · Flash-Lite · Embedding 2", GREEN)
    node(X0 + 6.8, Y0 + 1.05, 2.55, 0.95, "Firestore", "Catalog + vectors, inventory, signals, plans", MARIGOLD)
    node(X0 + 6.8, Y0 + 2.1, 2.55, 0.95, "BigQuery", "Sales history · AI.FORECAST (TimesFM)", GREEN)
    bottom = [("Model Armor", "Prompt screening", MAROON), ("Memory Bank", "Shopper preferences", GREEN),
              ("Google Search", "Trend grounding", MARIGOLD), ("Trace · Monitoring", "Spans, uptime, budget", MUTED)]
    for i, (t, b_, c) in enumerate(bottom):
        node(X0 + 1.85 + i * 1.9, Y0 + 3.2, 1.78, 0.85, t, b_, c, ts=9.5, bs=8)
    arrow(sl, X0 + 1.57, Y0 + 1.32, X0 + 1.83, Y0 + 1.32)
    arrow(sl, X0 + 3.72, Y0 + 1.32, X0 + 3.98, Y0 + 1.32)
    arrow(sl, X0 + 6.47, Y0 + 0.8, X0 + 6.78, Y0 + 0.45)
    arrow(sl, X0 + 6.47, Y0 + 1.4, X0 + 6.78, Y0 + 1.5)
    arrow(sl, X0 + 6.47, Y0 + 2.0, X0 + 6.78, Y0 + 2.55)
    arrow(sl, X0 + 5.2, Y0 + 2.37, X0 + 5.2, Y0 + 3.18)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    ap.add_argument("--links", type=Path, default=None, help="JSON with app/github/video URLs")
    a = ap.parse_args()
    links = json.loads(a.links.read_text()) if a.links else {}
    build(a.template, links)
